# -*- coding: utf-8 -*-
"""Ronda 3: navegación de semana en PLA-05 y chips con lunes siguiente en PLA-01.

Uso: python tools/update_inventory_docx3.py   (idempotente)
"""
import io, re, shutil
from docx import Document
from docx.shared import Pt

DOCX = r"C:\Users\Usuario\Desktop\Haymax\documentos\HAYMAX - Inventario de pantallas y navegacion.docx"
INV = r"C:\Users\Usuario\AppData\Local\Temp\opencode\inv.txt"
BACKUP = r"C:\Users\Usuario\AppData\Local\Temp\opencode\inventario.backup3.docx"

d = Document(DOCX)
done = []


def _rows(header_text, col):
    out = []
    for t in d.tables:
        hdr = [c.text.strip() for c in t.rows[0].cells]
        if len(hdr) > col and header_text in hdr[col]:
            out.append(t)
    return out


TCOMP = _rows("Componentes clave", 1)
TBOC = _rows("boceto", 2)


def _row(tables, sid):
    for t in tables:
        for r in t.rows[1:]:
            if r.cells[0].text.strip() == sid:
                return r
    raise KeyError(sid)


def add_bullet(sid, text):
    cell = _row(TCOMP, sid).cells[1]
    if text in cell.text:
        return
    p = cell.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(8.5)
    done.append("%s componentes += %s" % (sid, text[:50]))


def replace_in(tables, sid, col, old, new):
    cell = _row(tables, sid).cells[col]
    if new in cell.text:
        return
    if old not in cell.text:
        return
    for p in cell.paragraphs:
        runs = p.runs
        for r in runs:
            if old in r.text:
                r.text = r.text.replace(old, new)
                done.append("%s: %s -> %s" % (sid, old[:40], new[:40]))
                return
        if old in p.text:
            target = next((x for x in runs if x.font.size), runs[0])
            target.text = p.text.replace(old, new)
            for x in runs:
                if x is not target:
                    x.text = ""
            done.append("%s: %s -> %s" % (sid, old[:40], new[:40]))
            return


def append_to_cell(sid, extra):
    cell = _row(TBOC, sid).cells[2]
    if extra in cell.text:
        return
    p = cell.paragraphs[-1]
    r = p.add_run(extra)
    r.font.size = Pt(8.5)
    done.append("%s boceto += %s" % (sid, extra[:50]))


shutil.copyfile(DOCX, BACKUP)

# ---------- 1. Componentes ----------
replace_in(TCOMP, "PLA-01", 1,
           u"\u2022 Chips de d\u00eda (Hoy / Ma\u00f1ana / Vi / S\u00e1 / Do; la semana completa se edita en PLA-05)",
           u"\u2022 Chips de d\u00eda (Hoy / Ma\u00f1ana / S\u00e1bado / Pr\u00f3x. lunes: incluye el d\u00eda siguiente)")
add_bullet("PLA-05", u"\u2022 Cabecera de semana con navegaci\u00f3n visual \u2039 \u203a (lunes 18\u201324 nov)")
add_bullet("PLA-05", u"\u2022 Bloque \u00abPr\u00f3xima semana\u00bb con estado vac\u00edo y bot\u00f3n \u00abPlanificar un plato para esa semana\u00bb \u2192 PLA-01")

# ---------- 2. Boceto de PLA-05 ----------
append_to_cell("PLA-05", u"  \u00b7  Cabecera de semana \u2039 \u203a  \u00b7  Bloque Pr\u00f3xima semana (vac\u00edo) con CTA")

# ---------- 3. Decisiones de la seccion 9 ----------
NUEVAS = [
    u"PLA-05: la semana se cambia con \u00ab\u2039 \u203a\u00bb en la cabecera (visual en el prototipo) y la secci\u00f3n \u00abPr\u00f3xima semana\u00bb permite planificar con anticipaci\u00f3n: no hay que esperar al lunes. Los chips de PLA-01 incluyen \u00abPr\u00f3x. lunes\u00bb por la misma raz\u00f3n.",
    u"P\u00fablico objetivo (amas de casa, madres de edad avanzada): el plan usa d\u00edas relativos (Hoy / Ma\u00f1ana) en lugar de fechas, y cada pantalla tiene una sola acci\u00f3n principal con bot\u00f3n grande.",
]
size = None
for p in d.paragraphs:
    if p.style.name == "List Bullet" and p.runs:
        size = p.runs[0].font.size
        break
for t in NUEVAS:
    if any(t[:45] in p.text for p in d.paragraphs):
        continue
    p = d.add_paragraph(t, style="List Bullet")
    if p.runs and size:
        p.runs[0].font.size = size
    done.append("decision: " + t[:55])

d.save(DOCX)

# ---------- 4. volcar tablas para diffdoc.py ----------
doc2 = Document(DOCX)
lines = []
for i, t in enumerate(doc2.tables):
    for r in t.rows[1:]:
        cells = [c.text.replace("\n", " / ").strip() for c in r.cells]
        if cells and re.match(r"^[A-Z]{3}-\d{2}$", cells[0]):
            lines.append("T%d| %s" % (i, " | ".join(cells)))
io.open(INV, "w", encoding="utf-8").write("\n".join(lines))
io.open(r"C:\Users\Usuario\AppData\Local\Temp\opencode\docx_changes3.txt", "w",
        encoding="utf-8").write("\n".join(done))
print("OK: %d cambios ronda 3; backup %s" % (len(done), BACKUP))
