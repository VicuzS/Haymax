# -*- coding: utf-8 -*-
"""Ronda 4: modo invitado (datos locales), INI-01 priorizado y lenguaje sin
jerga (OCR -> escaneo, conflicto -> lenguaje cotidiano).

Uso: python tools/update_inventory_docx4.py   (idempotente)
"""
import io, re, shutil
from docx import Document
from docx.shared import Pt

DOCX = r"C:\Users\Usuario\Desktop\Haymax\documentos\HAYMAX - Inventario de pantallas y navegacion.docx"
INV = r"C:\Users\Usuario\AppData\Local\Temp\opencode\inv.txt"
BACKUP = r"C:\Users\Usuario\AppData\Local\Temp\opencode\inventario.backup4.docx"

d = Document(DOCX)
done = []


def _tables(header_text, col):
    out = []
    for t in d.tables:
        hdr = [c.text.strip() for c in t.rows[0].cells]
        if len(hdr) > col and header_text in hdr[col]:
            out.append(t)
    return out


TCOMP = _tables("Componentes clave", 1)
TBOC = _tables("boceto", 2)


def _row(tables, sid):
    for t in tables:
        for r in t.rows[1:]:
            if r.cells[0].text.strip() == sid:
                return r
    raise KeyError(sid)


def add_bullet(sid, text, tables=None):
    cell = _row(tables or TCOMP, sid).cells[1]
    if text in cell.text:
        return
    p = cell.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(8.5)
    done.append("%s componentes += %s" % (sid, text[:45]))


def append_to_cell(sid, extra, tables=None):
    cell = _row(tables or TBOC, sid).cells[2]
    if extra in cell.text:
        return
    p = cell.paragraphs[-1]
    r = p.add_run(extra)
    r.font.size = Pt(8.5)
    done.append("%s boceto += %s" % (sid, extra[:45]))


def _fix_paragraph(p, old, new):
    if old not in p.text:
        return False
    runs = p.runs
    for r in runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return True
    if runs:
        target = next((x for x in runs if x.font.size), runs[0])
        target.text = p.text.replace(old, new)
        for x in runs:
            if x is not target:
                x.text = ""
        return True
    return False


def replace_global(old, new):
    """Reemplaza una frase en TODAS las celdas y parrafos."""
    n = 0
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                while old in c.text and new not in c.text:
                    hit = False
                    for p in c.paragraphs:
                        if _fix_paragraph(p, old, new):
                            hit = True
                            n += 1
                    if not hit:
                        break
    for p in d.paragraphs:
        if _fix_paragraph(p, old, new):
            n += 1
    if n:
        done.append("%s -> %s (%d)" % (old[:40], new[:40], n))
    return n


shutil.copyfile(DOCX, BACKUP)

# ---------- 1. Modo invitado / cuenta opcional ----------
replace_global(
    u"Crea la cuenta o inicia sesión. Todos los casos de uso exigen \u201ccuenta activa\u201d como precondici\u00f3n, por lo que esta pantalla es obligatoria.",
    u"Crea la cuenta, inicia sesión o entra sin cuenta (modo local): la app se usa normalmente y, si luego te registras, lo guardado se sincroniza a la base de datos.")
replace_global(u"todas (precondición)", u"todas (cuenta opcional)")
replace_global(
    u"Acceso obligatorio: todos los casos de uso exigen cuenta activa.",
    u"Entrar sin cuenta (modo local) o crear cuenta; al registrarse se sincroniza lo guardado en el dispositivo.")
add_bullet("ACC-02", u"\u2022 Link Entrar sin cuenta (modo local; se sincroniza al registrarse)")
append_to_cell("ACC-02", u"  \u00b7  Link Entrar sin cuenta")

# ---------- 2. INI-01 priorizado ----------
add_bullet("INI-01", u"\u2022 Nota de vencimientos (alerta) junto a la despensa")
add_bullet("INI-01", u"\u2022 Secci\u00f3n Precios e insumos cerca de ti (\u2192 INI-03)")

# ---------- 3. PLA-04 en lenguaje cotidiano ----------
replace_global(
    u"Resuelve choques de horario: reemplazar el plato ya programado o agregar el nuevo de forma adicional.",
    u"Avisa con lenguaje simple cuando dos platos chocan: quitar el anterior o dejar los dos.")
replace_global(u"\u2022 Aviso de conflicto", u"\u2022 Aviso cotidiano (\u00abYa tienes Lomo Saltado el martes\u00bb)")
replace_global(u"\u2022 Opci\u00f3n Reemplazar (radio)", u"\u2022 Opci\u00f3n \u00abQuitar el plato anterior\u00bb (radio)")
replace_global(u"\u2022 Opci\u00f3n Agregar tambi\u00e9n (radio)", u"\u2022 Opci\u00f3n \u00abDejar los dos platos\u00bb (radio)")
replace_global(u"Reemplazar o agregar tambien, con radio buttons.",
               u"Quitar el plato anterior o dejar los dos, con radio buttons.")

# ---------- 4. Sin jerga: OCR -> escaneo ----------
replace_global(u"Revisi\u00f3n de datos OCR", u"Revisi\u00f3n del escaneo")
replace_global(u"c\u00e1mara (OCR)", u"c\u00e1mara (escaneo)")
replace_global(u"Primer paso del OCR", u"Primer paso del escaneo")
replace_global(u"Escanear boleta (OCR)", u"Escanear boleta (escaneo)")
replace_global(u"Escanear lista de compras (OCR)", u"Escanear lista de compras (escaneo)")
replace_global(u"plan de pago: OCR ilimitado", u"plan de pago: escaneo ilimitado")
replace_global(u"l\u00edmite de OCR", u"l\u00edmite de escaneo")
replace_global(u"\u2022 Link Probar el escaneo (OCR ilimitado en Premium)",
               u"\u2022 Link Probar el escaneo ilimitado (Premium)")
replace_global(u"Link Probar el escaneo (OCR)", u"Link Probar el escaneo ilimitado")
replace_global(u"OCR no reconoce nada", u"El escaneo no reconoce nada")
replace_global(u"Limite de OCR alcanzado", u"L\u00edmite de escaneo alcanzado")
replace_global(u"DES-05 (revisi\u00f3n OCR)", u"DES-05 (revisi\u00f3n del escaneo)")

# ---------- 5. Decisiones de la seccion 9 ----------
NUEVAS = [
    u"ACC-02 ofrece \u00abEntrar sin cuenta\u00bb: la app funciona con almacenamiento local (preferencias de los familiares y plan en el dispositivo); al crear cuenta esos datos se sincronizan a la base de datos. Evita el registro obligatorio para el p\u00fablico objetivo.",
    u"INI-01 reordenado y aligerado: arriba el plato de hoy, los accesos a planificar/ver plan y la alerta de vencimientos; debajo la despensa, el motor y el ahorro/salud. Se quit\u00f3 una foto duplicada del motor (la receta completa sigue en REC-01).",
    u"Lenguaje sin jerga para el p\u00fablico objetivo: \u00abOCR\u00bb se documenta como \u00abescaneo\u00bb y el paso de choque de horario (PLA-04) usa frases cotidianas (\u00abYa tienes Lomo Saltado el martes\u00bb / \u00abQuitar el plato anterior\u00bb / \u00abDejar los dos platos\u00bb).",
]
size = None
for p in d.paragraphs:
    if p.style.name == "List Bullet" and p.runs:
        size = p.runs[0].font.size
        break
for t in NUEVAS:
    if any(t[:45] in p.text for p in d.paragraphs):
        continue
    p = d.add_paragraph(t, style="List Bullet")
    if p.runs and size:
        p.runs[0].font.size = size
    done.append("decision: " + t[:55])

d.save(DOCX)

# ---------- 6. volcar tablas para diffdoc.py ----------
doc2 = Document(DOCX)
lines = []
for i, t in enumerate(doc2.tables):
    for r in t.rows[1:]:
        cells = [c.text.replace("\n", " / ").strip() for c in r.cells]
        if cells and re.match(r"^[A-Z]{3}-\d{2}$", cells[0]):
            lines.append("T%d| %s" % (i, " | ".join(cells)))
io.open(INV, "w", encoding="utf-8").write("\n".join(lines))
io.open(r"C:\Users\Usuario\AppData\Local\Temp\opencode\docx_changes4.txt", "w",
        encoding="utf-8").write("\n".join(done))
print("OK: %d cambios ronda 4; backup %s" % (len(done), BACKUP))
