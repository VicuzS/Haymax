# -*- coding: utf-8 -*-
"""Actualiza el DOCX de inventario para reflejar las correcciones de flujo
aplicadas al prototipo (tools/apply_flow_fixes.js).

Uso (ambos son idempotentes):
    python tools/update_inventory_docx.py
    python tools/update_inventory_docx2.py
"""
import io, os, re, shutil, sys
from docx import Document
from docx.shared import Pt

DOCX = r"C:\Users\Usuario\Desktop\Haymax\documentos\HAYMAX - Inventario de pantallas y navegacion.docx"
INV = r"C:\Users\Usuario\AppData\Local\Temp\opencode\inv.txt"
BACKUP = r"C:\Users\Usuario\AppData\Local\Temp\opencode\inventario.backup.docx"

d = Document(DOCX)
done = []


def row_of(ti, sid):
    for r in d.tables[ti].rows[1:]:
        if r.cells[0].text.strip() == sid:
            return r
    raise KeyError("T%d %s" % (ti, sid))


def add_dest(ti, sid, dest):
    """Agrega '-> XXX' a la celda Origen -> Destino (ultima parrafo)."""
    cell = row_of(ti, sid).cells[3]
    p = cell.paragraphs[-1]
    if dest in p.text:
        return
    run = p.add_run(u"\u2192 " + dest)
    run.font.size = Pt(8.5)
    done.append("%s destino += %s" % (sid, dest))


def add_bullet(ti, sid, text):
    cell = row_of(ti, sid).cells[1]
    if text in cell.text:
        return
    p = cell.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(8.5)
    done.append("%s componentes += %s" % (sid, text))


def replace_in(ti, sid, old, new, col=1):
    cell = row_of(ti, sid).cells[col]
    if new in cell.text:
        return
    if old not in cell.text:
        return
    for p in cell.paragraphs:
        for r in p.runs:
            if old in r.text:
                r.text = r.text.replace(old, new)
                done.append("%s: %s -> %s" % (sid, old[:30], new[:30]))
                return
    raise KeyError("no se encontro %r en T%d %s" % (old, ti, sid))


def replace_cell_text(ti, sid, old, new, col=1):
    cell = row_of(ti, sid).cells[col]
    if new in cell.text:
        return
    if old not in cell.text:
        return
    for p in cell.paragraphs:
        for r in p.runs:
            if old in r.text:
                r.text = r.text.replace(old, new)
                done.append("%s: %s -> %s" % (sid, old[:30], new[:30]))
                return
    # el texto puede estar partido en varios runs: reconstruir el parrafo
    for p in cell.paragraphs:
        if old in p.text:
            txt = p.text.replace(old, new)
            for r in list(p.runs)[1:]:
                r._element.getparent().remove(r._element)
            p.runs[0].text = txt
            done.append("%s: %s -> %s" % (sid, old[:30], new[:30]))
            return
    raise KeyError(old)


shutil.copyfile(DOCX, BACKUP)

# ---------- 1. Destinos (Origen -> Destino) que el prototipo ya tenia y el doc no ----------
add_dest(16, "PLA-01", "REC-02")     # tarjeta del plato abre la ficha
add_dest(12, "REC-01", "INI-03")     # Ver recetas por distrito
add_dest(20, "PRE-01", "INI-01")     # Omitir por ahora
add_dest(20, "PRE-03", "DES-01")     # Marcar como descartado
add_dest(16, "PLA-04", "PLA-01")     # Cancelar (demo) vuelve al paso 1
add_dest(24, "AJU-01", "ACC-04")     # Fila Mi distrito reutiliza ACC-04

# ACC-01: INI-01 es condicional (solo si habia sesion iniciada)
cell = row_of(8, "ACC-01").cells[3]
p = cell.paragraphs[-1]
if "condicional" not in p.text:
    r = p.add_run("  (INI-01 condicional: solo si habia sesion iniciada)")
    r.font.size = Pt(8.5)
    done.append("ACC-01 destino anotado como condicional")

# ---------- 2. Componentes clave ----------
add_bullet(9, "ACC-04", u"\u2022 Link \u00bfYa tienes cuenta? Inicia sesi\u00f3n")
add_bullet(9, "ACC-04", u"\u2022 Link Ya us\u00e9 Haymax antes \u00b7 ir a mi inicio")

add_bullet(11, "INI-01", u"\u2022 Link Ya lo prepar\u00e9 (descuenta insumos)")
add_bullet(11, "INI-01", u"\u2022 Link Ver mi plan de la semana (PLA-05)")
add_bullet(11, "INI-01", u"\u2022 Secci\u00f3n \u00bfQu\u00e9 puedo cocinar? (motor)")
add_bullet(11, "INI-01", u"\u2022 Tarjeta Balance de la semana")
add_bullet(11, "INI-01", u"\u2022 Chips Estado de despensa (disponibles / por vencer)")

add_bullet(13, "REC-01", u"\u2022 Link Falta 1 insumo \u2192 agregarlo a mi despensa")
add_bullet(13, "REC-01", u"\u2022 Link Planificar con lo que tengo")

replace_in(15, "DES-03", u"\u2022 Link \u00bfPor qu\u00e9 tres formas?",
           u"\u2022 Texto de ayuda \u00abTres formas seg\u00fan lo que tengas a mano\u00bb")
add_bullet(15, "DES-04", u"\u2022 Chips modo Boleta / Lista de compras")
add_bullet(15, "DES-04", u"\u2022 Nota de consejos de foto")
add_bullet(15, "DES-04", u"\u2022 Link Elegir otro tipo de documento")
add_bullet(15, "DES-05", u"\u2022 Link Repetir escaneo")
add_bullet(15, "DES-06", u"\u2022 Link Volver a las opciones de alta")

add_bullet(17, "PLA-01", u"\u2022 Chips de d\u00eda (Hoy / Ma\u00f1ana / Vi / S\u00e1 / Do; la semana completa se edita en PLA-05)")
add_bullet(17, "PLA-01", u"\u2022 Tarjeta del plato (abre la ficha REC-02)")
add_bullet(17, "PLA-03", u"\u2022 Link Ver mi lista de compras")

add_bullet(21, "PRE-02", u"\u2022 Link Volver a mi inicio")
add_bullet(21, "PRE-03", u"\u2022 Link Editar el \u00edtem desde la despensa")

add_bullet(25, "AJU-02", u"\u2022 Link Ver recetas que se ver\u00e1n afectadas")
replace_in(25, "AJU-04", u"\u2022 Link Restablecer compras",
           u"\u2022 Link Probar el escaneo (OCR ilimitado en Premium)")

# ---------- 3. Tabla "Componentes que se ven en el boceto" ----------
def append_to_cell(ti, sid, extra):
    cell = row_of(ti, sid).cells[2]
    if extra in cell.text:
        return
    p = cell.paragraphs[-1]
    r = p.add_run(extra)
    r.font.size = Pt(8.5)
    done.append("%s boceto += %s" % (sid, extra[:30]))

append_to_cell(26, "ACC-04", u"  \u00b7  Link \u00bfYa tienes cuenta? Inicia sesi\u00f3n  \u00b7  Link Ya us\u00e9 Haymax antes")
append_to_cell(27, "REC-01", u"  \u00b7  Link \u00abFalta 1 insumo\u00bb  \u00b7  Link \u00abPlanificar con lo que tengo\u00bb")
append_to_cell(28, "DES-05", u"  \u00b7  Link Repetir escaneo")
replace_cell_text(30, "AJU-04", u"Link Restablecer compras", u"Link Probar el escaneo (OCR)", col=2)

# ---------- 4. Flujo J3 (dia en PLA-01) ----------
cell = d.tables[6].rows[2].cells[2]
old, new = u"Elige porciones", u"Elige d\u00eda y porciones"
if old in cell.text and new not in cell.text:
    for p in cell.paragraphs:
        for r in p.runs:
            if old in r.text:
                r.text = r.text.replace(old, new)
                done.append("J3 paso 2: " + new)
                break

# ---------- 5. Seccion 9: pendientes resueltos + nuevas decisiones ----------
RESUELTOS = {
    u"Sincronizar los personajes":
        u"Resuelto: el prototipo unifica sobre Rosa y Diego (los textos de Figma ya no son fuente de verdad).",
    u"Crear la pantalla de revision de la lista de compras":
        u"Resuelto: DES-05 (revisi\u00f3n OCR) cubre boleta y lista; DES-04 ahora muestra el selector de modo Boleta / Lista de compras.",
    u"Agregar la alerta de vencimiento":
        u"Resuelto: PRE-03, AJU-02 y SAL-01 est\u00e1n construidas en el prototipo (31 pantallas).",
    u"Revisar la heur\u00edstica del bot\u00f3n flotante":
        u"Resuelto: el bot\u00f3n flotante de c\u00e1mara est\u00e1 en INI-01, REC-01 y DES-01, y navega a DES-03 (opciones de alta).",
    u"Renombrar los frames temporales":
        u"Resuelto: los frames del prototipo ya usan los nombres exactos de la secci\u00f3n 7.",
}
for p in d.paragraphs:
    for pref, nuevo in RESUELTOS.items():
        if p.text.startswith(pref):
            for r in list(p.runs)[1:]:
                r._element.getparent().remove(r._element)
            p.runs[0].text = nuevo
            done.append("pendiente resuelto: " + pref[:32])

NUEVAS = [
    u"ACC-01 \u2192 INI-01 es condicional (solo si hab\u00eda sesi\u00f3n iniciada): en el prototipo el splash encadena ACC-01 \u2192 ACC-02.",
    u"ACC-02: \u00abContinuar\u00bb y el ingreso con Google/Apple inician sesi\u00f3n (\u2192 INI-01); \u00ab\u00bfNo tienes cuenta? Reg\u00edstrate\u00bb abre el onboarding (\u2192 ACC-03).",
    u"ACC-04: \u00abEmpezar\u00bb \u2192 INI-02 (primera vez), \u00abYa us\u00e9 Haymax antes\u00bb \u2192 INI-01 e \u00ab\u00bfYa tienes cuenta?\u00bb \u2192 ACC-02.",
    u"PLA-01: se a\u00f1adi\u00f3 el selector de d\u00eda (chips Hoy / Ma\u00f1ana / fin de semana); el CU-06 y la HU-03 no especificaban el d\u00eda y el planificador lo necesita.",
    u"COM-01: \u00abAgregar \u00edtem manual\u00bb \u2192 DES-03 (opciones de alta) y \u00abRegistrar compra escaneada\u00bb \u2192 DES-04 (c\u00e1mara).",
    u"DES-03: la ayuda \u00ab\u00bfPor qu\u00e9 tres formas?\u00bb qued\u00f3 como texto inline (no existe una pantalla destino).",
    u"AJU-01: la fila \u00abMi distrito\u00bb reutiliza ACC-04 en lugar de crear una pantalla nueva.",
    u"Barra inferior: Inicio (INI-01), Despensa (DES-01), Planificar (PLA-05) y Ajustes (AJU-01); PRE-01 es modal, no pesta\u00f1a.",
    u"Salidas nuevas documentadas en el inventario: PLA-01 \u2192 REC-02, REC-01 \u2192 INI-03, PRE-01 \u2192 INI-01, PRE-03 \u2192 DES-01, PLA-04 \u2192 PLA-01, AJU-01 \u2192 ACC-04.",
]
size = None
for p in d.paragraphs:
    if p.style.name == "List Bullet" and p.runs:
        size = p.runs[0].font.size
        break
for t in NUEVAS:
    if any(t[:40] in p.text for p in d.paragraphs):
        continue
    p = d.add_paragraph(t, style="List Bullet")
    if p.runs and size:
        p.runs[0].font.size = size
    done.append("nueva decision: " + t[:44])

d.save(DOCX)

# ---------- 6. Volcar las tablas a inv.txt para diffdoc.py ----------
doc2 = Document(DOCX)
lines = []
for i, t in enumerate(doc2.tables):
    for r in t.rows[1:]:
        cells = [c.text.replace("\n", " / ").strip() for c in r.cells]
        if cells and re.match(r"^[A-Z]{3}-\d{2}$", cells[0]):
            lines.append("T%d| %s" % (i, " | ".join(cells)))
io.open(INV, "w", encoding="utf-8").write("\n".join(lines))

io.open(r"C:\Users\Usuario\AppData\Local\Temp\opencode\docx_changes.txt", "w",
        encoding="utf-8").write("\n".join(done))
print("OK: %d cambios; backup en %s" % (len(done), BACKUP))
