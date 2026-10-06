# -*- coding: utf-8 -*-
"""Ronda 6: accesos rapidos a la lista de compras (COM-01).
   A) REC-02: link de faltantes -> COM-01
   B) DES-01: link 'Lo que falta' -> COM-01
   C) INI-01: la tarjeta Lista de compras queda junto al plato de hoy

Uso: python tools/update_inventory_docx6.py   (idempotente)
"""
import io, re, shutil
from docx import Document
from docx.shared import Pt

DOCX = r"C:\Users\Usuario\Desktop\Haymax\documentos\HAYMAX - Inventario de pantallas y navegacion.docx"
INV = r"C:\Users\Usuario\AppData\Local\Temp\opencode\inv.txt"
BACKUP = r"C:\Users\Usuario\AppData\Local\Temp\opencode\inventario.backup6.docx"

d = Document(DOCX)
done = []


def _fix_paragraph(p, old, new):
    if old not in p.text:
        return False
    runs = p.runs
    for r in runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return True
    if runs:
        target = next((x for x in runs if x.font.size), runs[0])
        target.text = p.text.replace(old, new)
        for x in runs:
            if x is not target:
                x.text = ""
        return True
    return False


def replace_global(old, new):
    n = 0
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                while old in c.text and new not in c.text:
                    hit = False
                    for p in c.paragraphs:
                        if _fix_paragraph(p, old, new):
                            hit = True
                            n += 1
                    if not hit:
                        break
    for p in d.paragraphs:
        if _fix_paragraph(p, old, new):
            n += 1
    if n:
        done.append("%s -> %s (%d)" % (old[:45], new[:45], n))
    return n


def _tables(header_text, col):
    out = []
    for t in d.tables:
        hdr = [c.text.strip() for c in t.rows[0].cells]
        if len(hdr) > col and header_text in hdr[col]:
            out.append(t)
    return out


TCOMP = _tables("Componentes clave", 1)


def _row(tables, sid):
    for t in tables:
        for r in t.rows[1:]:
            if r.cells[0].text.strip() == sid:
                return r
    raise KeyError(sid)


def add_bullet(sid, text):
    cell = _row(TCOMP, sid).cells[1]
    if text in cell.text:
        return
    p = cell.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(8.5)
    done.append("%s componentes += %s" % (sid, text[:45]))


def add_destino(sid, target):
    """Agrega un destino al parrafo que empieza con flecha de la fila dada."""
    arrow = u"\u2192 %s" % target
    for t in d.tables:
        for r in t.rows[1:]:
            if r.cells[0].text.strip() != sid:
                continue
            for c in r.cells:
                for p in c.paragraphs:
                    if p.text.strip().startswith(u"\u2192"):
                        if arrow in p.text:
                            done.append("%s destinos ya incluye %s" % (sid, target))
                            return
                        if not _fix_paragraph(p, p.text, p.text + arrow):
                            continue
                        done.append("%s destinos += %s" % (sid, target))
                        return
    raise KeyError("destinos de %s" % sid)


shutil.copyfile(DOCX, BACKUP)

# C) INI-01: tarjeta de lista junto al plato de hoy
replace_global(u"\u2022 Tarjeta progreso Lista de compras",
               u"\u2022 Tarjeta progreso Lista de compras (junto al plato de hoy, "
               u"acceso r\u00e1pido)")

# A/B) componentes y destinos de los nuevos atajos
add_destino("REC-02", "COM-01")
add_destino("DES-01", "COM-01")
add_bullet("REC-02", u"\u2022 Link Agregar faltantes a mi lista de compras (\u2192 COM-01)")
add_bullet("DES-01", u"\u2022 Link Lo que falta est\u00e1 en tu lista de compras (\u2192 COM-01)")

# Decision de la seccion 9
NUEVA = (u"La lista de compras (COM-01) queda a dos toques desde cualquier punto "
         u"del recorrido: tarjeta junto al plato de hoy en INI-01, atajo de "
         u"faltantes en el detalle de receta (REC-02), atajo desde la despensa "
         u"(DES-01) y bot\u00f3n de generar/ver lista en PLA-03 y PLA-05.")
if not any(NUEVA[:45] in p.text for p in d.paragraphs):
    size = None
    for p in d.paragraphs:
        if p.style.name == "List Bullet" and p.runs:
            size = p.runs[0].font.size
            break
    p = d.add_paragraph(NUEVA, style="List Bullet")
    if p.runs and size:
        p.runs[0].font.size = size
    done.append("decision: " + NUEVA[:55])

d.save(DOCX)

# volcar tablas para diffdoc.py
doc2 = Document(DOCX)
lines = []
for i, t in enumerate(doc2.tables):
    for r in t.rows[1:]:
        cells = [c.text.replace("\n", " / ").strip() for c in r.cells]
        if cells and re.match(r"^[A-Z]{3}-\d{2}$", cells[0]):
            lines.append("T%d| %s" % (i, " | ".join(cells)))
io.open(INV, "w", encoding="utf-8").write("\n".join(lines))
io.open(r"C:\Users\Usuario\AppData\Local\Temp\opencode\docx_changes6.txt", "w",
        encoding="utf-8").write("\n".join(done))
print("OK: %d cambios ronda 6; backup %s" % (len(done), BACKUP))
