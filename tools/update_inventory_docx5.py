# -*- coding: utf-8 -*-
"""Ronda 5: INI-01 — la sugerencia de comida se lista sin imagen
(la foto se ve al entrar a la receta, REC-02).

Uso: python tools/update_inventory_docx5.py   (idempotente)
"""
import io, re, shutil
from docx import Document
from docx.shared import Pt

DOCX = r"C:\Users\Usuario\Desktop\Haymax\documentos\HAYMAX - Inventario de pantallas y navegacion.docx"
INV = r"C:\Users\Usuario\AppData\Local\Temp\opencode\inv.txt"
BACKUP = r"C:\Users\Usuario\AppData\Local\Temp\opencode\inventario.backup5.docx"

d = Document(DOCX)
done = []


def _tables(header_text, col):
    out = []
    for t in d.tables:
        hdr = [c.text.strip() for c in t.rows[0].cells]
        if len(hdr) > col and header_text in hdr[col]:
            out.append(t)
    return out


TCOMP = _tables("Componentes clave", 1)


def _row(tables, sid):
    for t in tables:
        for r in t.rows[1:]:
            if r.cells[0].text.strip() == sid:
                return r
    raise KeyError(sid)


def _fix_paragraph(p, old, new):
    if old not in p.text:
        return False
    runs = p.runs
    for r in runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return True
    if runs:
        target = next((x for x in runs if x.font.size), runs[0])
        target.text = p.text.replace(old, new)
        for x in runs:
            if x is not target:
                x.text = ""
        return True
    return False


def replace_global(old, new):
    n = 0
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                while old in c.text and new not in c.text:
                    hit = False
                    for p in c.paragraphs:
                        if _fix_paragraph(p, old, new):
                            hit = True
                            n += 1
                    if not hit:
                        break
    for p in d.paragraphs:
        if _fix_paragraph(p, old, new):
            n += 1
    if n:
        done.append("%s -> %s (%d)" % (old[:45], new[:45], n))
    return n


shutil.copyfile(DOCX, BACKUP)

# 1. Componente de INI-01: sugerencia en tarjeta compacta sin foto
replace_global(u"\u2022 Secci\u00f3n Sugerencias con precio por porci\u00f3n",
               u"\u2022 Secci\u00f3n Sugerencias: tarjeta compacta con precio por "
               u"porci\u00f3n (sin foto; la foto est\u00e1 en REC-02)")

# 2. Decision de la seccion 9
NUEVA = (u"INI-01 lista las comidas con los datos principales (porciones, tiempo "
         u"y precio) y sin imagen: la foto aparece reci\u00e9n al entrar a la receta "
         u"(REC-02). Asi la pantalla principal queda ligera y no compite con la "
         u"alerta de vencimientos ni con el plato del dia.")
if not any(NUEVA[:45] in p.text for p in d.paragraphs):
    size = None
    for p in d.paragraphs:
        if p.style.name == "List Bullet" and p.runs:
            size = p.runs[0].font.size
            break
    p = d.add_paragraph(NUEVA, style="List Bullet")
    if p.runs and size:
        p.runs[0].font.size = size
    done.append("decision: " + NUEVA[:55])

d.save(DOCX)

# 3. volcar tablas para diffdoc.py
doc2 = Document(DOCX)
lines = []
for i, t in enumerate(doc2.tables):
    for r in t.rows[1:]:
        cells = [c.text.replace("\n", " / ").strip() for c in r.cells]
        if cells and re.match(r"^[A-Z]{3}-\d{2}$", cells[0]):
            lines.append("T%d| %s" % (i, " | ".join(cells)))
io.open(INV, "w", encoding="utf-8").write("\n".join(lines))
io.open(r"C:\Users\Usuario\AppData\Local\Temp\opencode\docx_changes5.txt", "w",
        encoding="utf-8").write("\n".join(done))
print("OK: %d cambios ronda 5; backup %s" % (len(done), BACKUP))
