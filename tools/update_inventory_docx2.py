# -*- coding: utf-8 -*-
"""Ronda 2 de correcciones del DOCX de inventario: destinos nuevos, columna
Origen corregida (pestaña / FAB / inconsistencias), componentes y decisiones.

Uso: python tools/update_inventory_docx2.py   (idempotente)
"""
import io, re, shutil
from docx import Document
from docx.shared import Pt

DOCX = r"C:\Users\Usuario\Desktop\Haymax\documentos\HAYMAX - Inventario de pantallas y navegacion.docx"
INV = r"C:\Users\Usuario\AppData\Local\Temp\opencode\inv.txt"
BACKUP = r"C:\Users\Usuario\AppData\Local\Temp\opencode\inventario.backup2.docx"

d = Document(DOCX)
done = []


def _table(header_key, col):
    for t in d.tables:
        hdr = [c.text.strip() for c in t.rows[0].cells]
        if len(hdr) > col and header_key in hdr[col]:
            return t
    raise KeyError(header_key)


TDEST = _table("Origen", 3)      # no basta: hay varias tablas; buscamos por fila
TDESTS = [t for t in d.tables
          if len(t.rows[0].cells) >= 4 and "Origen" in t.rows[0].cells[3].text]
TCOMP = [t for t in d.tables
         if len(t.rows[0].cells) >= 4 and t.rows[0].cells[1].text.strip() == "Componentes clave"]
TBOC = [t for t in d.tables
        if len(t.rows[0].cells) >= 3 and "boceto" in t.rows[0].cells[2].text]


def _row(tables, sid):
    for t in tables:
        for r in t.rows[1:]:
            if r.cells[0].text.strip() == sid:
                return r
    raise KeyError(sid)


def add_dest(sid, dest):
    r = _row(TDESTS, sid)
    p = r.cells[3].paragraphs[-1]
    if dest in p.text:
        return
    run = p.add_run(u"\u2192 " + dest)
    run.font.size = Pt(8.5)
    done.append("%s destino += %s" % (sid, dest))


def set_origins(sid, text):
    p = _row(TDESTS, sid).cells[3].paragraphs[0]
    if p.text.strip() == text:
        return
    runs = p.runs                      # wrappers: mantener la MISMA lista
    target = next((x for x in runs if x.font.size), runs[0])
    target.text = text
    for x in runs:
        if x is not target:
            x.text = ""
    done.append("%s origenes = %s" % (sid, text[:60]))


def add_bullet(sid, text):
    cell = _row(TCOMP, sid).cells[1]
    if text in cell.text:
        return
    p = cell.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(8.5)
    done.append("%s componentes += %s" % (sid, text[:45]))


def replace_in(tables, sid, col, old, new):
    cell = _row(tables, sid).cells[col]
    if new in cell.text:
        return
    if old not in cell.text:
        return
    for p in cell.paragraphs:
        for r in p.runs:
            if old in r.text:
                r.text = r.text.replace(old, new)
                done.append("%s: %s -> %s" % (sid, old[:35], new[:35]))
                return
    for p in cell.paragraphs:
        if old in p.text:
            runs = p.runs
            target = next((x for x in runs if x.font.size), runs[0])
            target.text = p.text.replace(old, new)
            for x in runs:
                if x is not target:
                    x.text = ""
            done.append("%s: %s -> %s" % (sid, old[:35], new[:35]))
            return


def replace_anywhere(old, new):
    n = 0
    for t in d.tables:
        for r in t.rows[1:]:
            for c in r.cells:
                while old in c.text and new not in c.text:
                    hit = False
                    for p in c.paragraphs:
                        for run in p.runs:
                            if old in run.text:
                                run.text = run.text.replace(old, new)
                                hit = True
                                break
                        if hit:
                            break
                    if not hit:
                        # el texto esta repartido entre runs: reescribe el run con tamaño
                        p = next(p for p in c.paragraphs if old in p.text)
                        runs = p.runs
                        target = next((x for x in runs if x.font.size), runs[0])
                        target.text = p.text.replace(old, new)
                        for x in runs:
                            if x is not target:
                                x.text = ""
                    n += 1
    if n:
        done.append("tabla: %s -> %s (%d)" % (old[:35], new[:35], n))
    return n


shutil.copyfile(DOCX, BACKUP)

# ---------------- 1. Destinos que ahora si existen en el prototipo ----------------
add_dest("ACC-04", "AJU-03")   # + Agregar integrante reutiliza Integrantes del hogar
add_dest("INI-01", "PLA-05")   # Ver mi plan de la semana
add_dest("PLA-01", "INI-01")   # Link Cancelar
add_dest("PLA-03", "PLA-01")   # Link Cancelar
add_dest("PLA-05", "PLA-01")   # Editar por tarjeta

# ---------------- 2. Columna Origen: sin contradicciones ----------------
set_origins("INI-01", "ACC-02, ACC-04, PLA-02, PRE-02  \u00b7  PLA-05, SAL-01, COM-01, DES-08 (por pesta\u00f1a)")
set_origins("REC-01", "INI-01")
set_origins("DES-01", "INI-01, DES-08, PRE-02  \u00b7  REC-01 (por pesta\u00f1a)")
set_origins("DES-03", "INI-02, COM-01  \u00b7  DES-01 e INI-01 (por bot\u00f3n flotante)")
set_origins("PLA-05", "PLA-02, PLA-03, PLA-04  \u00b7  INI-01 (por pesta\u00f1a)")
set_origins("AJU-01", "Tab Ajustes")
set_origins("AJU-03", "AJU-01, ACC-04")

# ---------------- 3. Componentes clave ----------------
replace_in(TCOMP, "INI-01", 1, u"\u2022 Link Editar mi plan semanal",
           u"\u2022 Link Ver mi plan de la semana (PLA-05)")
add_bullet("INI-01", u"\u2022 Link Planificar un plato nuevo (PLA-01)")
add_bullet("PLA-01", u"\u2022 Links de demo: resultado con faltantes y conflicto de horario")
replace_in(TCOMP, "PLA-01", 1, u"\u2022 Chips de d\u00eda (Hoy / Ma\u00f1ana / fin de semana)",
           u"\u2022 Chips de d\u00eda (Hoy / Ma\u00f1ana / Vi / S\u00e1 / Do; la semana completa se edita en PLA-05)")
replace_in(TCOMP, "DES-01", 1, u"\u2022 Items con foto, cantidad y estado",
           u"\u2022 Items con icono, cantidad, estado y badge de vencimiento naranja")
replace_in(TCOMP, "AJU-01", 1, u"\u2022 Perfil y correo",
           u"\u2022 Perfil (Rosa Ticlla \u2192 integrantes del hogar)")
replace_in(TCOMP, "AJU-01", 1, u"\u2022 Fila Notificaciones",
           u"\u2022 Fila Notificaciones (estado \u00abActivadas\u00bb; el interruptor es de la app)")

# ---------------- 4. CTA de DES-08 renombrado en todo el documento ----------------
replace_anywhere("Continuar a mi plan", "Continuar y planificar")

# ---------------- 5. Decisiones nuevas (seccion 9) ----------------
NUEVAS = [
    u"ACC-01: el splash ahora tiene adem\u00e1s el link \u00ab\u00bfYa ten\u00edas sesi\u00f3n? Entra a mi inicio\u00bb \u2192 INI-01 (antes el destino condicional no exist\u00eda en el prototipo).",
    u"ACC-04: \u00ab+ Agregar integrante\u00bb abre AJU-03 (Integrantes del hogar): no existe una pantalla de alta de integrantes en las 31, as\u00ed no queda un bot\u00f3n sin destino.",
    u"INI-01 separa \u00abVer mi plan de la semana\u00bb (PLA-05) de \u00abPlanificar un plato nuevo\u00bb (PLA-01): antes un \u00fanico link dec\u00eda \u00abeditar mi plan\u00bb pero abr\u00eda el paso 1 de porciones.",
    u"PLA-01 y PLA-03 ahora s\u00ed tienen \u00abLink Cancelar\u00bb clicable (PLA-01 \u2192 INI-01, PLA-03 \u2192 PLA-01); se quit\u00f3 el \u00abCancelar\u00bb decorativo de la cabecera de PLA-01 (no era clicable).",
    u"PLA-05: cada tarjeta muestra el estado (Completa / Falta N insumos) y su link \u00abEditar\u00bb \u2192 PLA-01, m\u00e1s el resumen de cobertura de la semana.",
    u"DES-01: badge naranja de vencimiento en Leche, Palta y Cilantro (los mismos 3 de PRE-03 y de la nota).",
    u"AJU-01: \u00abRosa Ticlla\u00bb \u2192 AJU-03; \u00abNotificaciones\u00bb queda como fila con estado \u00abActivadas\u00bb (el interruptor real pertenece a la app, no al prototipo est\u00e1tico).",
    u"DES-08: el CTA \u00abContinuar a mi plan\u00bb se renombr\u00f3 \u00abContinuar y planificar\u00bb porque su destino es PLA-01 (paso 1), no PLA-05.",
    u"Columna Origen corregida: se marcan los destinos cubiertos por pesta\u00f1a o bot\u00f3n flotante y se quitan los inexistentes (INI-02 \u2192 REC-01, PRE-03 \u2192 REC-01, DES-08 \u2192 PLA-05, ACC-04 \u2192 AJU-01).",
    u"Los chips de d\u00eda (PLA-01) y de la semana (PLA-05) son visuales en el prototipo est\u00e1tico: el d\u00eda elegido no se propaga; en la app s\u00ed se reflejar\u00e1 en PLA-05 y en la lista de compras.",
]
size = None
for p in d.paragraphs:
    if p.style.name == "List Bullet" and p.runs:
        size = p.runs[0].font.size
        break
for t in NUEVAS:
    if any(t[:45] in p.text for p in d.paragraphs):
        continue
    p = d.add_paragraph(t, style="List Bullet")
    if p.runs and size:
        p.runs[0].font.size = size
    done.append("decision: " + t[:50])

d.save(DOCX)

# ---------------- 6. volcar tablas para diffdoc.py ----------------
doc2 = Document(DOCX)
lines = []
for i, t in enumerate(doc2.tables):
    for r in t.rows[1:]:
        cells = [c.text.replace("\n", " / ").strip() for c in r.cells]
        if cells and re.match(r"^[A-Z]{3}-\d{2}$", cells[0]):
            lines.append("T%d| %s" % (i, " | ".join(cells)))
io.open(INV, "w", encoding="utf-8").write("\n".join(lines))
io.open(r"C:\Users\Usuario\AppData\Local\Temp\opencode\docx_changes2.txt", "w",
        encoding="utf-8").write("\n".join(done))
print("OK: %d cambios ronda 2; backup %s" % (len(done), BACKUP))
